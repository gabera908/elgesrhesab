"""التصدير — Excel و CSV للتقارير."""
import io
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.reports import _opening_balances, _period_balances
from app.core.deps import require_permission
from app.database import get_db
from app.models.account import Account
from app.models.user import User

router = APIRouter(prefix="/api/export", tags=["التصدير"])

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def _trial_balance_rows(db: Session, from_date: date, to_date: date):
    accounts = db.scalars(select(Account).where(Account.is_active == True)).all()  # noqa: E712
    opening = _opening_balances(db, from_date)
    period = _period_balances(db, from_date, to_date)

    rows = []
    for acc in accounts:
        op_d, op_c = opening.get(acc.id, (0, 0))
        pr_d, pr_c = period.get(acc.id, (0, 0))
        if not (op_d or op_c or pr_d or pr_c):
            continue
        rows.append(
            {
                "code": acc.code,
                "name": acc.name,
                "opening_debit": float(op_d),
                "opening_credit": float(op_c),
                "period_debit": float(pr_d),
                "period_credit": float(pr_c),
                "closing_debit": float(op_d + pr_d),
                "closing_credit": float(op_c + pr_c),
            }
        )
    return rows


@router.get("/trial-balance/csv")
async def export_trial_balance_csv(
    from_date: date = Query(...),
    to_date: date = Query(...),
    current_user: User = Depends(require_permission("reports", "read")),
    db: Session = Depends(get_db),
):
    """تصدير ميزان المراجعة CSV."""
    import csv

    rows = _trial_balance_rows(db, from_date, to_date)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        ["الكود", "الحساب", "افتتاحي مدين", "افتتاحي دائن",
         "حركة مدين", "حركة دائن", "ختامي مدين", "ختامي دائن"]
    )
    for r in rows:
        writer.writerow(
            [r["code"], r["name"], r["opening_debit"], r["opening_credit"],
             r["period_debit"], r["period_credit"], r["closing_debit"], r["closing_credit"]]
        )

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=trial-balance-{from_date}-{to_date}.csv"},
    )


@router.get("/trial-balance/excel")
async def export_trial_balance_excel(
    from_date: date = Query(...),
    to_date: date = Query(...),
    current_user: User = Depends(require_permission("reports", "read")),
    db: Session = Depends(get_db),
):
    """تصدير ميزان المراجعة Excel منسّق."""
    if not HAS_OPENPYXL:
        from fastapi import HTTPException
        raise HTTPException(status_code=501, detail="مكتبة Excel غير مثبتة")

    rows = _trial_balance_rows(db, from_date, to_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ميزان المراجعة"
    ws.sheet_view.rightToLeft = True

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="D97757", end_color="D97757", fill_type="solid")
    thin = Side(style="thin", color="E5E3DA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # العنوان
    ws.merge_cells("A1:H1")
    ws["A1"] = f"ميزان المراجعة — {from_date} إلى {to_date}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["الكود", "الحساب", "افتتاحي مدين", "افتتاحي دائن",
               "حركة مدين", "حركة دائن", "ختامي مدين", "ختامي دائن"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for r_idx, r in enumerate(rows, 4):
        values = [r["code"], r["name"], r["opening_debit"], r["opening_credit"],
                  r["period_debit"], r["period_credit"], r["closing_debit"], r["closing_credit"]]
        for c_idx, v in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.border = border
            if c_idx > 2:
                cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="center" if c_idx > 2 else "right")

    # صف الإجمالي
    total_row = len(rows) + 4
    ws.cell(row=total_row, column=1, value="الإجمالي").font = Font(bold=True)
    for c in range(3, 9):
        col_letter = openpyxl.utils.get_column_letter(c)
        cell = ws.cell(row=total_row, column=c)
        cell.value = f"=SUM({col_letter}4:{col_letter}{total_row - 1})"
        cell.font = Font(bold=True)
        cell.number_format = "#,##0.00"
        cell.border = border

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    ws.column_dimensions["B"].width = 35

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=trial-balance-{from_date}-{to_date}.xlsx"},
    )
